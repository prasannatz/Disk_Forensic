from pptx import Presentation
from pptx.util import Inches

prs = Presentation(r'c:\Users\lenovo\Desktop\disk_forensic\05_Presentation\MP08_Disk_Forensics_Completed.pptx')
print(f'Slide dimensions: {prs.slide_width / Inches(1):.2f}" x {prs.slide_height / Inches(1):.2f}"')

for idx in [4, 5, 6]:
    slide = prs.slides[idx]
    print(f'=== Slide {idx+1} ===')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if '[' in p.text and 'screenshot' in p.text.lower():
                    print(f'Placeholder found: "{p.text}"')
                    print(f'Shape pos: left={shape.left/Inches(1):.2f}", top={shape.top/Inches(1):.2f}", width={shape.width/Inches(1):.2f}", height={shape.height/Inches(1):.2f}"')

